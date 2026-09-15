import io
import json
import re
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from quiz import export, quizfile, widgets
from quiz.samples import sample_quiz
from quiz.scoring import rank
from quiz.store import QuizStore

from helpers import answer_all, make_store


def test_quiz_file_round_trip_keeps_questions_answers_and_settings_but_no_secrets():
    store, clock, code, key = make_store(winners=5, email_domains=["example.com"])
    store.open(code)
    person = store.join(code, "Player One", "player.one@example.com")
    answer_all(store, clock, code, person.token)
    quiz = store.get(code)
    data = quizfile.dump(quiz)
    text = data.decode("utf-8")
    assert key not in text and code not in text and "player.one" not in text.lower()
    title, description, questions, settings = quizfile.load(data)
    copy, _ = QuizStore().create(title, description, questions, settings)
    assert [q.text for q in copy.questions] == [q.text for q in quiz.questions]
    assert [[o.text for o in q.options] for q in copy.questions] == [[o.text for o in q.options] for q in quiz.questions]
    assert [len(q.correct) for q in copy.questions] == [len(q.correct) for q in quiz.questions]
    assert copy.questions[2].accepted == ["Mars"]
    assert copy.settings.winners == 5 and copy.settings.email_domains == ["example.com"]
    assert not copy.issues()


@pytest.mark.parametrize("data, message", [
    pytest.param(b"not json", "valid JSON", id="not-json"),
    pytest.param(b"[]", "saved by this app", id="list"),
    pytest.param(json.dumps({"kind": "something-else"}).encode(), "saved by this app", id="other-kind"),
    pytest.param(json.dumps({"kind": "meeting-quiz"}).encode(), "no questions", id="no-questions"),
    pytest.param(json.dumps({"kind": "meeting-quiz", "questions": [{"type": "essay"}]}).encode(), "unknown type",
                 id="unknown-type"),
    pytest.param(None, "too big", id="too-big"),
])
def test_invalid_quiz_files_are_rejected(data, message):
    with pytest.raises(quizfile.QuizFileError, match=message):
        quizfile.load(b" " * (quizfile.MAX_FILE_BYTES + 1) if data is None else data)


def test_hostile_values_in_quiz_files_are_shaped_and_limited():
    payload = {
        "kind": "meeting-quiz", "title": ["not", "text"],
        "settings": {"winners": float("inf"), "ask_email": "yes", "not_eligible": [1, "a@example.com"]},
        "questions": [{"type": "single", "text": "x" * 5000, "points": 10 ** 9,
                       "options": [{"text": str(i), "right": True} for i in range(40)]}] * 80,
    }
    title, _, questions, settings = quizfile.load(json.dumps(payload).replace("Infinity", "1e999").encode())
    quiz, _ = QuizStore().create(title, "", questions, settings)
    assert quiz.title == "Untitled quiz"
    assert len(quiz.questions) == 50
    assert len(quiz.questions[0].options) == 10 and len(quiz.questions[0].correct) == 1
    assert len(quiz.questions[0].text) == 500 and quiz.questions[0].points == 1000
    assert quiz.settings.ask_email is True and quiz.settings.not_eligible == ["a@example.com"]


def test_sample_quiz_is_ready_to_open():
    store = QuizStore()
    quiz, _ = store.create(*sample_quiz())
    assert not quiz.issues()
    store.open(quiz.code)


def test_results_workbook_has_every_sheet_and_neutralizes_formulas():
    store, clock, code, _ = make_store()
    store.open(code)
    person = store.join(code, "=HYPERLINK(\"https://example.com\")", "formula@example.com")
    answer_all(store, clock, code, person.token)
    quiz = store.get(code)
    results = rank(quiz)
    workbook = load_workbook(io.BytesIO(export.results_workbook(quiz, results, clock(), "America/Sao_Paulo")))
    assert workbook.sheetnames == ["Summary", "Ranking", "Answers", "Questions"]
    ranking = workbook["Ranking"]
    headers = [cell.value for cell in ranking[1]]
    name = ranking.cell(row=2, column=headers.index("Name") + 1).value
    assert name.startswith("'=")
    answers = workbook["Answers"]
    assert "Q3 answer" in [cell.value for cell in answers[1]]
    csv = export.to_csv(export.ranking_frame(results)).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv


def test_ranking_table_can_hide_emails_for_screen_sharing():
    store, clock, code, _ = make_store()
    store.open(code)
    person = store.join(code, "Player One", "player.one@example.com")
    answer_all(store, clock, code, person.token)
    frame = export.ranking_frame(rank(store.get(code)), include_email=False)
    assert "Email" not in frame.columns
    assert "player.one@example.com" not in frame.to_string()


def test_widgets_never_let_names_break_out_of_the_script():
    store, clock, code, _ = make_store()
    store.open(code)
    person = store.join(code, "</script><img src=x onerror=alert(1)>&", "x@example.com")
    answer_all(store, clock, code, person.token)
    quiz = store.get(code)
    page = widgets.winners_html(quiz, rank(quiz))
    script = page.split("<script>", 1)[1]
    assert "</script><img" not in script and "<img" not in page
    assert "\\u003c/script\\u003e" in script
    assert "__PAYLOAD__" not in page
    countdown = widgets.countdown_html(30, 60, large=True)
    assert "__PAYLOAD__" not in countdown and '"remaining": 30' in countdown


@pytest.mark.skipif(shutil.which("node") is None, reason="node is not installed")
@pytest.mark.parametrize("template", ["winners.html", "countdown.html"])
def test_widget_scripts_are_valid_javascript(tmp_path, template):
    quiz, _ = QuizStore().create(*sample_quiz())
    page = widgets.winners_html(quiz, rank(quiz)) if template == "winners.html" else widgets.countdown_html(5, 10)
    script = re.search(r"<script>(.*)</script>", page, re.S).group(1)
    path = tmp_path / "widget.js"
    path.write_text(script, encoding="utf-8")
    result = subprocess.run(["node", "--check", str(path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
